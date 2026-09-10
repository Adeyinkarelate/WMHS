"""Build docs/chapter-4.docx — Chapter Four: System Implementation and Results."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
DIAG = ROOT / "diagrams"
OUT = ROOT / "chapter-4.docx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
INK = RGBColor(0x1E, 0x29, 0x3B)


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "94A3B8")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def para_format(p, first_indent=True, after=8, align="justify", space=1.5):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if space == 1.5 else WD_LINE_SPACING.SINGLE
    if first_indent:
        pf.first_line_indent = Cm(1.25)
    else:
        pf.first_line_indent = Cm(0)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]


def add_body(doc, text, first_indent=True):
    p = doc.add_paragraph()
    para_format(p, first_indent=first_indent)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_heading_styled(doc, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(16 if level == 2 else 12)
    pf.space_after = Pt(8)
    pf.first_line_indent = Cm(0)
    pf.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    size = 14 if level == 2 else 13
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=NAVY)
    return p


def add_chapter_title(doc):
    for line, size, space_before in (
        ("CHAPTER FOUR", 16, 24),
        ("SYSTEM IMPLEMENTATION AND RESULTS", 16, 6),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(line)
        set_run_font(run, size=size, bold=True, color=NAVY)


def add_caption(doc, text, kind="Figure"):
    p = doc.add_paragraph()
    para_format(p, first_indent=False, after=14, align="center", space=1)
    run = p.add_run(text)
    set_run_font(run, size=11, italic=True, color=NAVY)


def add_figure(doc, filename, caption, width=6.2):
    path = DIAG / filename
    if not path.exists():
        add_body(doc, f"[Missing figure file: {filename}]", first_indent=False)
        add_caption(doc, caption)
        return
    p = doc.add_paragraph()
    para_format(p, first_indent=False, after=4, align="center", space=1)
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=12)


def add_table(doc, caption, headers, rows, col_widths=None):
    add_caption(doc, caption, "Table")
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(cell, "0B1F3A")
        set_cell_border(cell)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10)
            if r_i % 2 == 0:
                shade_cell(cell, "F1F5F9")
            set_cell_border(cell)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = INK

    add_chapter_title(doc)

    # 4.1
    add_heading_styled(doc, "4.1  Introduction", 2)
    add_body(
        doc,
        "This chapter presents the implementation of the Web Maternal Health System (WMHS), "
        "a web-based maternal health monitoring, risk-assessment, and referral application "
        "designed for the Nigerian antenatal setting. It describes the programming languages, "
        "development tools, technology stack, and implementation environment, and it shows how "
        "the three-tier design specified in Chapter Three was realised. The chapter then "
        "documents the clinical risk engine, security controls, functional verification of the "
        "implemented workflows, and the deployed user interfaces.",
    )
    add_body(
        doc,
        "Implementation followed the Chapter Three specifications: a presentation tier that "
        "provides the patient and provider portals; an application tier that hosts REST APIs, "
        "authentication, risk scoring, and referral matching; and a data tier that stores users, "
        "clinical records, assessments, alerts, and referrals. The implemented system is "
        "decision support. It does not diagnose disease, replace a clinician, or dispatch "
        "emergency transport.",
    )

    # 4.2
    add_heading_styled(doc, "4.2  Choice of Programming Language and Tools", 2)
    add_body(
        doc,
        "The technologies were selected for reliability, maintainability, and suitability for a "
        "cloud-hosted health information system that can be used from ordinary web browsers in "
        "Nigerian facilities and homes. Preference was given to a single full-stack TypeScript "
        "codebase so that the same types describe forms, API payloads, and database records.",
    )

    add_heading_styled(doc, "4.2.1  Programming Languages", 3)
    add_body(
        doc,
        "TypeScript is the primary language for both the user interface and the server-side API. "
        "Static type checking reduces integration errors between screens and endpoints, and it "
        "documents clinical fields such as risk level, score, and user role in the source code. "
        "JavaScript (ES6+) is the runtime language produced from TypeScript. Asynchronous request "
        "handling is appropriate for concurrent logins, symptom submissions, and risk recalculation.",
    )
    add_body(
        doc,
        "Python is used only for an optional, offline Random Forest training script "
        "(ml/train_model.py). Production inference does not load Python or a serialised "
        ".joblib artefact, so the Vercel deployment does not depend on a Python runtime. "
        "SQL is generated by Prisma against PostgreSQL in production (Neon) and against SQLite "
        "during local development.",
    )

    add_heading_styled(doc, "4.2.2  Development Tools and Frameworks", 3)
    add_body(doc, "Frontend (presentation tier)", first_indent=False)
    add_bullets(
        doc,
        [
            "Next.js 14 (App Router) and React 18 render the marketing site, authentication pages, patient portal, and provider dashboard. Server Components load caseload and risk data on the server; Client Components handle interactive forms and charts.",
            "Tailwind CSS provides a consistent, responsive clinical layout.",
            "Framer Motion supplies page transitions.",
            "Recharts draws the provider risk-distribution chart.",
        ],
    )
    add_body(doc, "Backend (application tier)", first_indent=False)
    add_bullets(
        doc,
        [
            "Node.js is the runtime for Next.js.",
            "Next.js Route Handlers (app/api/...) implement REST endpoints. There is no separate Express.js server.",
            "NextAuth.js (Credentials provider, JWT session, eight-hour expiry) handles login. Passwords are compared with bcryptjs.",
            "Zod validates registration, login, symptoms, tests, profile updates, referral status, and alert status.",
            "Prisma 5 is the object-relational mapper for queries, relations, and schema management.",
        ],
    )
    add_body(doc, "Database (data tier)", first_indent=False)
    add_bullets(
        doc,
        [
            "PostgreSQL hosted on Neon is the production store, with connection pooling, backups, and separate environments.",
            "Local development may use SQLite through the DATABASE_URL environment variable.",
        ],
    )
    add_body(doc, "Clinical decision support", first_indent=False)
    add_bullets(
        doc,
        [
            "Production scoring is implemented in TypeScript (lib/ai/riskService.ts): a weighted ensemble aligned with WHO-, ACOG- and NICE-style thresholds, plus rule-based emergency overrides and complication screens for pre-eclampsia, gestational diabetes, and preterm labour.",
            "An optional scikit-learn Random Forest trainer is documented in ml/README.md but is not invoked at request time.",
        ],
    )
    add_body(
        doc,
        "Supporting tools included Visual Studio Code, Git and GitHub, npm, environment-variable "
        "configuration (dotenv / Vercel), Prisma Studio for inspecting tables, and Vercel for "
        "application hosting.",
    )

    # 4.3
    add_heading_styled(doc, "4.3  Implementation Environment", 2)
    add_heading_styled(doc, "4.3.1  Hardware Requirements", 3)
    add_body(
        doc,
        "Development and testing were carried out on a workstation equipped with an Intel Core "
        "i7-10750H processor, 16 GB RAM, and a 512 GB solid-state drive, running Windows 11 Pro "
        "(64-bit). Production uses cloud infrastructure: the Next.js application is deployed on "
        "Vercel, and the PostgreSQL database is hosted on Neon. End users require only a current "
        "web browser and an internet connection. No specialised clinical hardware is required on "
        "the client.",
    )

    add_heading_styled(doc, "4.3.2  Software Requirements", 3)
    add_table(
        doc,
        "Table 4.1: Software requirements for development and production",
        ["Category", "Software"],
        [
            ["Operating system (development)", "Windows 11 Pro (64-bit)"],
            ["Runtime", "Node.js 20+"],
            ["Framework", "Next.js 14, React 18, TypeScript 5"],
            ["Styling and charts", "Tailwind CSS, Recharts, Framer Motion"],
            ["Authentication", "NextAuth.js 4, bcryptjs"],
            ["Validation", "Zod"],
            ["ORM and database", "Prisma 5, PostgreSQL 16 (Neon); SQLite (local)"],
            ["Editor and version control", "Visual Studio Code, Git, GitHub"],
            ["Hosting", "Vercel (application), Neon (database)"],
            ["Security in transit", "HTTPS / TLS via the hosting platform"],
        ],
    )

    # 4.4
    add_heading_styled(doc, "4.4  System Architecture Implementation", 2)
    add_body(
        doc,
        "The system was implemented using the three-tier architecture described in Chapter Three. "
        "The presentation tier comprises React pages for patients (/patient/...), providers "
        "(/provider/...), and public authentication and marketing routes. The application tier "
        "comprises Next.js API routes for registration, session management, symptoms, tests, risk "
        "assessment, facilities, referrals, notifications, and the health assistant, together with "
        "NextAuth, role guards, the clinical risk engine, and Haversine facility search. The data "
        "tier comprises Prisma models mapped to User, Patient, Provider, Facility, Symptom, "
        "TestResult, RiskAssessment, Alert, Referral, Notification, PatientAssignment, and "
        "PasswordResetToken.",
    )
    add_body(
        doc,
        "Actors reach the system through a web browser over HTTPS. The presentation tier sends "
        "JSON REST requests to the application tier. The application tier reads and writes the "
        "database through Prisma. Figure 4.1 shows the implemented architecture.",
    )
    add_figure(
        doc,
        "system-architecture.png",
        "Figure 4.1: System architecture of WMHS showing presentation, application, and data tiers",
        width=6.3,
    )

    add_heading_styled(doc, "4.4.1  Frontend Implementation", 3)
    add_body(
        doc,
        "The frontend is organised as Next.js route groups. Public routes include the marketing "
        "home page, login, registration, and password reset. The patient portal provides a "
        "dashboard, symptom log, test entry and upload, risk results, history, referrals, a "
        "contextual health assistant, and profile management. The provider portal provides a "
        "caseload dashboard, patient list and detail, alerts, and referral management.",
    )
    add_body(
        doc,
        "The patient dashboard displays the latest risk level and score, gestational age derived "
        "from the last menstrual period using Naegele’s rule, parity, active alerts, recent "
        "symptoms and tests, and open referrals. The provider dashboard lists assigned patients "
        "with their latest risk, highlights high-risk cases, shows active alerts and pending "
        "referrals, and includes a risk-distribution chart (low, medium, and high). Role-based "
        "routing is enforced in middleware.ts: unauthenticated users are redirected to /login; "
        "patients cannot open provider routes and providers cannot open patient routes.",
    )
    add_figure(
        doc,
        "fig-patient-dashboard.png",
        "Figure 4.2: Patient portal dashboard showing risk indicator, gestational age, and recent records",
        width=6.3,
    )
    add_body(
        doc,
        "Symptom submission is implemented as POST /api/symptoms. After Zod validation, the "
        "symptom is stored and the risk engine is executed on the patient’s latest symptoms and "
        "tests. The handler returns the new symptom, the stored assessment, and the engine result "
        "so that the client can navigate to the risk-results page.",
    )

    add_heading_styled(doc, "4.4.2  Backend Implementation", 3)
    add_body(
        doc,
        "The backend is a set of Next.js Route Handlers rather than a standalone Express "
        "application. The principal modules are described below.",
    )
    add_body(doc, "Authentication module", first_indent=False)
    add_bullets(
        doc,
        [
            "POST /api/auth/register creates a User. For patients it also creates a Patient profile, computes the estimated date of delivery from the last menstrual period, optionally assigns the first available provider, and stores default Lagos coordinates for nearby-facility search. For providers it creates a Provider profile linked to a facility. Passwords are hashed with bcryptjs using 12 salt rounds. Registration is rate-limited (eight attempts per IP address per minute).",
            "POST /api/auth/[...nextauth] performs NextAuth credentials login and issues a JWT session containing id, role, patientId, and providerId.",
            "POST /api/auth/forgot-password issues a 30-minute password-reset token and returns a generic message to avoid email enumeration. The endpoint is rate-limited.",
            "Route guards (requireUser and assertPatientAccess) return HTTP 401 or 403 and restrict providers to assigned patients.",
        ],
    )
    add_body(doc, "Patient and clinical-record modules", first_indent=False)
    add_bullets(
        doc,
        [
            "Patient profile retrieval and update.",
            "Symptom create, read, update, and delete.",
            "Test-result create, read, update, and delete.",
            "Risk-assessment history for a given patient.",
        ],
    )
    add_body(doc, "AI analysis (risk) module", first_indent=False)
    add_bullets(
        doc,
        [
            "POST /api/risk/assess and automatic invocation after symptom writes.",
            "runAndStoreRisk extracts features, scores risk, stores a RiskAssessment row, opens Alert rows for triggered complications, and notifies the patient and assigned providers when the level is HIGH.",
        ],
    )
    add_body(doc, "Referral module", first_indent=False)
    add_bullets(
        doc,
        [
            "Facility list and detail, and GET /api/facilities/nearby using the Haversine formula (lib/utils/haversine.ts), with an optional filter for emergency-capable sites.",
            "Referral create and status updates (PENDING, APPROVED, COMPLETED, CANCELLED), with a patient notification when a referral is created.",
        ],
    )
    add_body(doc, "Notifications and health assistant", first_indent=False)
    add_bullets(
        doc,
        [
            "In-application notifications for high-risk assessments and new referrals.",
            "POST /api/assistant returns knowledge-based answers using the patient’s latest risk and gestational age, with escalation flags for danger-sign questions.",
        ],
    )
    add_figure(
        doc,
        "fig-api-architecture.png",
        "Figure 4.3: API architecture showing presentation clients, authentication, clinical endpoints, and the data layer",
        width=6.3,
    )

    add_heading_styled(doc, "4.4.3  Database Implementation", 3)
    add_body(
        doc,
        "The PostgreSQL schema was implemented from the Chapter Three design using "
        "prisma/schema.prisma. Identifiers are universally unique identifiers (UUIDs). Clinical "
        "fields that are lists or structured objects in the application (risk factors, "
        "recommendations, services offered, and pre-existing conditions) are stored as text and "
        "serialised in application code. Figure 4.4 shows the entity-relationship diagram of the "
        "implemented database.",
    )
    add_figure(
        doc,
        "er-diagram.png",
        "Figure 4.4: Entity-relationship diagram of the WMHS database",
        width=6.3,
    )
    add_body(
        doc,
        "Schema changes are applied with Prisma (prisma db push or prisma migrate). The "
        "updatedAt columns are maintained by Prisma. Secondary indexes exist on foreign keys, "
        "role, riskLevel, assessedAt, referral status, and alert status and severity.",
    )
    add_table(
        doc,
        "Table 4.2: Principal data relationships implemented in Prisma",
        ["Relationship", "Cardinality", "Mechanism"],
        [
            ["User–Patient; User–Provider", "1 : 0..1", "Exclusive by role"],
            ["User–Notification; User–PasswordResetToken", "1 : N", "Cascade delete"],
            ["Patient–Symptom, TestResult, RiskAssessment, Alert, Referral", "1 : N", "Cascade delete"],
            ["Patient–Provider", "N : M", "PatientAssignment (unique pair)"],
            ["Facility–Provider", "1 : 0..N", "Optional facilityId"],
            ["Facility–Referral", "1 : N", "Receiving facility"],
        ],
    )
    add_body(
        doc,
        "There is no separate one-to-one “AI recommendation profile” table. Recommendations are "
        "stored on each RiskAssessment row. Figure 4.5 shows representative records from the "
        "Patient, Symptom, and RiskAssessment tables, consistent with the seeded demonstration data.",
    )
    add_figure(
        doc,
        "fig-sample-records.png",
        "Figure 4.5: Sample records from the Patient, Symptom, and RiskAssessment tables",
        width=6.3,
    )

    # 4.5
    add_heading_styled(doc, "4.5  AI Model Implementation", 2)
    add_body(
        doc,
        "Production inference is the TypeScript clinical engine in lib/ai/riskService.ts. The "
        "engine is a weighted ensemble that approximates a Random Forest-style combination of "
        "vital-sign and danger-sign features. Coefficients are aligned with published clinical "
        "thresholds rather than a private labelled hospital dataset. A Python Random Forest may "
        "be trained offline; that artefact is not loaded at runtime on Vercel.",
    )

    add_heading_styled(doc, "4.5.1  Feature Extraction", 3)
    add_body(
        doc,
        "The function extractFeaturesFromRecords constructs fourteen features from maternal age, "
        "the latest matching test results, and recent symptoms whose severity is at least 2. "
        "Missing vitals fall back to physiologically plausible defaults so that a complete feature "
        "vector is always available. After scoring, contextual modifiers may be added: gestational "
        "age from LMP, body-mass index from height and weight, and grand multiparity when parity "
        "is five or more.",
    )
    add_table(
        doc,
        "Table 4.3: Features used by the clinical risk engine",
        ["Feature", "Typical source"],
        [
            ["Age", "Date of birth"],
            ["Systolic blood pressure", "Latest matching TestResult (default 110 mmHg)"],
            ["Diastolic blood pressure", "Latest matching TestResult (default 70 mmHg)"],
            ["Blood glucose", "Latest matching TestResult (default 90 mg/dL)"],
            ["Heart rate", "Latest matching TestResult (default 80 bpm)"],
            ["Body temperature", "Latest matching TestResult (default 36.8 °C)"],
            ["Severe headache; blurred vision; swelling", "Symptom type and severity"],
            ["Vaginal bleeding; abdominal pain; dizziness", "Symptom type and severity"],
            ["Reduced fetal movement; high stress", "Symptom type and severity"],
        ],
    )

    add_heading_styled(doc, "4.5.2  Scoring and Classification", 3)
    add_body(
        doc,
        "Each feature contributes a weighted partial score. The sum is clamped to the interval "
        "[0, 1]. Classification uses both the numeric score and hard emergency rules. The level "
        "is HIGH if the score is at least 0.55, or if any of the following are present: vaginal "
        "bleeding; reduced fetal movement; systolic blood pressure of 160 mmHg or more together "
        "with diastolic pressure of 110 mmHg or more; or severe headache combined with blurred "
        "vision. The level is MEDIUM if the score is at least 0.28, and LOW otherwise. Supporting "
        "calculations include Naegele’s estimated date of delivery, gestational age in weeks, "
        "mean arterial pressure, and body-mass index.",
    )

    add_heading_styled(doc, "4.5.3  Complication Screening and Recommendations", 3)
    add_body(
        doc,
        "When indicated, the engine screens for pre-eclampsia, gestational diabetes, and preterm "
        "labour, and may add an emergency screen. The function buildRecommendations returns "
        "suggested tests, actions, a short care plan, and an urgency band (monitor, consult, or "
        "emergency). A HIGH result creates alerts and in-application notifications for the patient "
        "and assigned providers. Figure 4.6 summarises the pipeline from record collection to "
        "persistence.",
    )
    add_figure(
        doc,
        "fig-risk-pipeline.png",
        "Figure 4.6: Clinical risk-assessment pipeline from feature extraction to alerts",
        width=5.8,
    )

    add_heading_styled(doc, "4.5.4  Optional Random Forest Training", 3)
    add_body(
        doc,
        "The script ml/train_model.py can fit a scikit-learn RandomForestClassifier (200 trees, "
        "maximum depth 8, balanced class weights) on synthetic data (2,000 rows, 80/20 train–test "
        "split) and write model.joblib. That file is an artefact of the product specification and "
        "is not part of the production request path. Any hold-out accuracy printed by the script "
        "applies only to the synthetic generator and is not reported here as clinical performance.",
    )

    # 4.6
    add_heading_styled(doc, "4.6  Security Implementation", 2)
    add_heading_styled(doc, "4.6.1  Authentication and Authorisation", 3)
    add_body(
        doc,
        "Users submit an email address and password on /login. Zod validates the payload. "
        "NextAuth looks up the user and compares the password with the stored bcrypt hash. On "
        "success a JSON Web Token is issued with an eight-hour lifetime and stored in an "
        "HTTP-only cookie. Middleware protects /patient/* and /provider/* routes. API handlers "
        "call requireUser (optionally with an allowed-role list). Providers may access only "
        "patients to whom they are assigned (assertPatientAccess). Figure 4.7 shows the flow.",
    )
    add_figure(
        doc,
        "fig-auth-flow.png",
        "Figure 4.7: Authentication and authorisation flow",
        width=6.2,
    )

    add_heading_styled(doc, "4.6.2  Data Security Measures", 3)
    add_bullets(
        doc,
        [
            "Password hashing: bcryptjs with 12 salt rounds. Only passwordHash is stored on the User table.",
            "HTTPS/TLS: enforced by Vercel in production.",
            "Input validation: Zod schemas on authentication and clinical payloads, which reduces injection and malformed-data risk.",
            "Rate limiting: an in-memory limiter on registration and password-reset endpoints to reduce brute-force attempts.",
            "Role-based access control: JWT role, edge middleware, and API guards.",
            "Session control: eight-hour JWT; logout through NextAuth.",
            "Password reset: single-use tokens that expire after 30 minutes; generic API responses.",
            "Least privilege on records: patients see only their own data; providers see their assigned caseload.",
        ],
    )

    # 4.7
    add_heading_styled(doc, "4.7  Results of Testing", 2)
    add_body(
        doc,
        "Verification used functional (black-box) tests against the implemented routes and user "
        "interfaces, with the seeded demonstration accounts. The repository does not contain an "
        "automated Jest or Mocha suite; therefore unit-coverage percentages are not claimed. "
        "The tests below confirm that the implemented modules behave as specified in Chapter Three.",
    )

    add_heading_styled(doc, "4.7.1  Functional Testing Results", 3)
    add_table(
        doc,
        "Table 4.4: Functional test results",
        ["S/N", "Module", "Test case", "Expected result", "Result"],
        [
            ["1", "Registration", "Valid patient data", "User and Patient created; EDD from LMP", "Pass"],
            ["2", "Registration", "Duplicate email", "HTTP 409", "Pass"],
            ["3", "Registration", "Invalid password", "Validation error", "Pass"],
            ["4", "Login", "Valid credentials", "Session cookie; role dashboard", "Pass"],
            ["5", "Login", "Wrong password", "Sign-in rejected", "Pass"],
            ["6", "Authorisation", "Patient opens provider area", "Redirect away from /provider", "Pass"],
            ["7", "Authorisation", "Unauthenticated patient route", "Redirect to /login", "Pass"],
            ["8", "Symptoms", "Log a danger-sign symptom", "Symptom stored; risk recalculated", "Pass"],
            ["9", "Tests", "Enter blood pressure / glucose", "TestResult stored", "Pass"],
            ["10", "Risk", "High-risk vitals or symptoms", "HIGH with factors and recommendations", "Pass"],
            ["11", "Alerts", "Triggered complication", "Alert created if none active", "Pass"],
            ["12", "Notifications", "HIGH assessment", "Patient and providers notified", "Pass"],
            ["13", "Facilities", "Known latitude/longitude", "Facilities sorted by Haversine km", "Pass"],
            ["14", "Referral", "Create referral", "PENDING row; patient notified", "Pass"],
            ["15", "Referral", "Provider updates status", "Status persisted", "Pass"],
            ["16", "Assistant", "Danger-sign question", "Answer with escalate flag", "Pass"],
            ["17", "Password reset", "Known email", "Token created", "Pass"],
            ["18", "Profile", "Update phone / LMP", "Patient profile updated", "Pass"],
        ],
    )
    add_body(
        doc,
        "All eighteen functional cases passed on the seeded system. Failures of validation and "
        "authorisation were confirmed as rejected requests rather than silent success.",
    )

    add_heading_styled(doc, "4.7.2  Integration Testing Results", 3)
    add_body(
        doc,
        "Integration tests exercised complete workflows that cross the presentation, application, "
        "and data tiers.",
    )
    add_table(
        doc,
        "Table 4.5: End-to-end workflow results",
        ["Workflow", "Steps", "Result"],
        [
            ["Patient monitoring", "Register → login → log symptom → view risk", "Pass"],
            ["Vitals and risk", "Enter SBP/DBP → assess → dashboard updates", "Pass"],
            ["Provider caseload", "Provider login → list by risk → open patient", "Pass"],
            ["Alert handling", "High-risk case → alert on dashboard → resolve", "Pass"],
            ["Referral", "Nearby facilities → create referral → track status", "Pass"],
            ["Cross-role isolation", "Patient cannot read another patient’s API data", "Pass"],
        ],
    )

    add_heading_styled(doc, "4.7.3  Risk-Engine Verification", 3)
    add_body(
        doc,
        "The engine was checked with scenario cases that correspond to the implemented rules. "
        "These checks confirm correct classification behaviour; they are not a claim of accuracy "
        "on an external labelled clinical dataset.",
    )
    add_table(
        doc,
        "Table 4.6: Risk-engine scenario checks",
        ["Scenario", "Illustrative inputs", "Expected class", "Result"],
        [
            ["Routine antenatal visit", "Age 28; BP 110/70; glucose 90; no danger signs", "LOW", "Pass"],
            ["Raised blood pressure", "SBP 142; DBP 92; headache", "MEDIUM or HIGH", "Pass"],
            ["Emergency blood pressure", "SBP 165; DBP 112", "HIGH", "Pass"],
            ["Combined neurological signs", "Severe headache and blurred vision", "HIGH", "Pass"],
            ["Antepartum bleeding", "Vaginal bleeding logged", "HIGH", "Pass"],
            ["Reduced fetal movement", "Low fetal movement logged", "HIGH", "Pass"],
            ["Hyperglycaemia", "Glucose ≥ 140 mg/dL", "Elevated score; GDM screen if HIGH", "Pass"],
        ],
    )

    add_heading_styled(doc, "4.7.4  Performance Observations", 3)
    add_body(
        doc,
        "The production stack is serverless (Vercel) with hosted PostgreSQL (Neon). Informal "
        "checks on the development workstation (npm run dev) showed interactive page loads and "
        "API responses completing in under about two seconds for dashboard rendering, symptom "
        "submission with risk persistence, and nearby-facility sorting over the seeded facility "
        "list. These observations indicate that the implemented operations are suitable for "
        "interactive clinical use at demonstration scale. Formal load testing with a published "
        "tool is recommended before large-facility roll-out.",
    )
    add_table(
        doc,
        "Table 4.7: Qualitative performance on the development workstation",
        ["Operation", "Observation"],
        [
            ["Login and dashboard", "Acceptable for interactive use"],
            ["Symptom submit and risk persist", "Completes in one request cycle"],
            ["Nearby facility sort", "In-process Haversine over seeded facilities"],
            ["Provider caseload and chart", "Acceptable for the seeded caseload"],
        ],
    )

    add_heading_styled(doc, "4.7.5  Usability Review", 3)
    add_body(
        doc,
        "A heuristic review of the implemented interfaces (labels, risk colour coding, empty "
        "states, and mobile-width Tailwind layout) found that the main tasks are completable: "
        "register, log a symptom, read a risk result, open referrals, and review a provider "
        "caseload. Formal field usability testing with antenatal clients and midwives was not "
        "conducted as a measured survey for this chapter and is recommended as future work. "
        "During implementation, copy was kept free of diagnostic claims, tap targets were kept "
        "large, and high-risk language was made visually prominent.",
    )

    # 4.8
    add_heading_styled(doc, "4.8  System Features Demonstration", 2)
    add_heading_styled(doc, "4.8.1  Patient Registration Flow", 3)
    add_body(
        doc,
        "Figure 4.8 shows the patient registration interface. The user opens /register, chooses "
        "the patient role, and enters identity and obstetric data, including last menstrual "
        "period. The server hashes the password, creates the User and Patient rows, computes the "
        "estimated date of delivery, and may assign a provider. The user then signs in and is "
        "redirected to /patient/dashboard.",
    )
    add_figure(
        doc,
        "fig-registration.png",
        "Figure 4.8: Patient registration interface",
        width=5.4,
    )
    add_bullets(
        doc,
        [
            "Step 1: The user navigates to the registration page.",
            "Step 2: Personal and obstetric information is completed and the terms are accepted.",
            "Step 3: An account is created with a hashed password.",
            "Step 4: The user is signed in and redirected to the patient dashboard.",
        ],
    )

    add_heading_styled(doc, "4.8.2  Symptom Submission and AI Analysis", 3)
    add_body(
        doc,
        "Figure 4.9 shows the symptom form and the risk-results page. The patient selects a "
        "common pregnancy symptom and a severity from 1 to 5, optionally adds notes, and submits "
        "the form. The API stores the symptom, runs the risk engine, and the interface displays "
        "the level, score, contributing factors, care plan, and suggested tests or actions.",
    )
    add_figure(
        doc,
        "fig-symptom-risk.png",
        "Figure 4.9: Symptom submission form and risk-result display",
        width=6.3,
    )

    add_heading_styled(doc, "4.8.3  Provider Dashboard", 3)
    add_body(
        doc,
        "Figure 4.10 shows the provider caseload dashboard. High-risk patients are emphasised "
        "through counts, colour, and an alert feed. A risk-distribution chart summarises the "
        "caseload. From this screen a provider can open a patient record, review alerts, and "
        "manage referrals.",
    )
    add_figure(
        doc,
        "fig-provider-dashboard.png",
        "Figure 4.10: Provider dashboard showing caseload statistics, risk distribution, and alerts",
        width=6.3,
    )
    add_bullets(
        doc,
        [
            "Risk-based patient list: high-risk cases are highlighted with colour coding.",
            "Real-time alerts: a feed of active and critical cases.",
            "Health trends: graphical distribution of low, medium, and high risk.",
            "Referral management: pending referrals are counted and linked.",
        ],
    )

    add_heading_styled(doc, "4.8.4  Referral System", 3)
    add_body(
        doc,
        "Figure 4.11 illustrates referral recommendations. Nearby facilities are ordered by "
        "Haversine distance from the patient’s stored coordinates (defaulting to Lagos if unset). "
        "Emergency-capable sites can be filtered. Each facility shows services, contact "
        "information, and capacity. Created referrals are tracked through PENDING, APPROVED, "
        "COMPLETED, and CANCELLED.",
    )
    add_figure(
        doc,
        "fig-referral.png",
        "Figure 4.11: Referral interface with existing referral status and nearby facilities",
        width=6.2,
    )

    # 4.9
    add_heading_styled(doc, "4.9  Deployment", 2)
    add_heading_styled(doc, "4.9.1  Frontend and Application Deployment", 3)
    add_body(
        doc,
        "The Next.js application, including both the user interface and the API routes, is "
        "deployed on Vercel. The repository is connected to Vercel; production environment "
        "variables are set in the Vercel dashboard (DATABASE_URL, NEXTAUTH_SECRET, NEXTAUTH_URL, "
        "and related secrets); and each push to the main branch triggers next build. Prisma "
        "Client is generated during install (postinstall). In production the Prisma datasource "
        "provider is PostgreSQL when pointing at Neon. A custom domain may be mapped if required.",
    )

    add_heading_styled(doc, "4.9.2  Database Deployment", 3)
    add_body(
        doc,
        "The PostgreSQL database is deployed on Neon (PostgreSQL 16). After the project is "
        "created, the schema is applied with Prisma and demonstration users and facilities are "
        "loaded with prisma db seed. Connection pooling is used for serverless scalability. Neon "
        "provides scheduled backups and point-in-time recovery according to the project plan. "
        "Figure 4.12 shows the deployment arrangement.",
    )
    add_figure(
        doc,
        "fig-deployment.png",
        "Figure 4.12: Deployment architecture on Vercel and Neon",
        width=6.2,
    )

    # 4.10
    add_heading_styled(doc, "4.10  Summary", 2)
    add_body(
        doc,
        "This chapter described the implementation of WMHS as a single Next.js application with "
        "a three-tier architecture, Prisma-backed storage, NextAuth security, and an on-server "
        "clinical risk engine with Haversine referral matching. The presentation tier provides "
        "separate patient and provider portals. The application tier implements authentication, "
        "clinical recording, risk scoring, alerts, notifications, and referrals. The data tier "
        "persists identity, maternal records, assessments, and care-coordination data on "
        "PostgreSQL.",
    )
    add_body(
        doc,
        "Functional and workflow tests of the main paths passed on the seeded system. Scenario "
        "checks confirmed that emergency danger signs raise a HIGH classification. Screenshots "
        "and architecture figures document the realised design. The system remains decision "
        "support for antenatal monitoring and referral; it is not a substitute for emergency "
        "obstetric care.",
    )

    doc.save(OUT)
    print("wrote", OUT, "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    build()
